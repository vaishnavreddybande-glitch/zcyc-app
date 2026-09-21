"""Logic behind streamlit_app.py. Nothing here imports Streamlit, so it can be tested and reused on its own.

Curves are fitted exactly as in the submission (engine.bootstrap = exact fit, engine.ls_fit = smoothed fit);
this module only wires them to the interface: data loading, tables, STRIPS pricing, file upload, figures.
"""
import glob
import io
import json
import os
import re
from dataclasses import dataclass
from datetime import date, datetime
from functools import lru_cache

import numpy as np
import pandas as pd
import openpyxl
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from scipy.interpolate import CubicSpline

import engine as en
import multidate as md
import strips as sp

ROOT = md.ROOT
MULTI = md.MULTI
UPLOADED = md.UPLOADED
REQUIRED_TBILLS = ["7 Days", "3 Months", "6 Months", "12 Months"]


def base(iso):
    """The calendar date of a date key. Dates entered by a visitor of a shared deployment are stored as 'YYYY-MM-DD~<session id>'."""
    return iso.split("~")[0]


def multiuser(path=None):
    """True when the app serves several visitors at once (Streamlit Community Cloud, or ZCYC_MULTIUSER=1), so that every visitor's
    uploads must stay private. On a laptop it is False and entered dates persist in input_data/uploaded/."""
    v = os.environ.get("ZCYC_MULTIUSER")
    if v is not None:
        return v.strip().lower() in ("1", "true", "yes")
    return os.path.abspath(path or __file__).replace("\\", "/").startswith("/mount/src")


def set_upload_dir(path):
    """Where entered dates are stored (a temporary folder on a shared deployment)."""
    global UPLOADED
    UPLOADED = path
    md.UPLOADED = path


def purge_old_uploads(hours=6):
    """Delete entered files older than `hours` (shared deployments only)."""
    import time
    if not os.path.isdir(UPLOADED):
        return
    cutoff = time.time() - hours * 3600
    for f in glob.glob(os.path.join(UPLOADED, "*")):
        try:
            if os.path.getmtime(f) < cutoff:
                os.remove(f)
        except OSError:
            pass
OUT = os.path.join(ROOT, "additional_material", "outputs")
HELD_OUT = "2026-09-11"          # the only date not used to choose the smoothing settings
NAVY, ORANGE, TEAL, GREY, RED = "#1F3A5F", "#E07B39", "#2A9D8F", "#8A8F98", "#C0392B"
SEGS = [(0, 1), (1, 3), (3, 5), (5, 10), (10, 14), (14, 20), (20, 30), (30, 40), (40, 50)]
METHODS = {"exact": ("Exact-fit bootstrap", GREY, ":"), "smooth": ("Smoothed fit", ORANGE, "--")}


def tuned_params():
    """lambda and slope selected on the four August dates by tune_and_checks.py (falls back to the README values)."""
    try:
        p = json.load(open(os.path.join(OUT, "tuned_params.json")))
        return float(p["lam"]), float(p["slope"])
    except Exception:
        return 0.005, 0.5


# ------------------------------------------------------------------ data
def _dates_in(folder):
    out = set()
    for f in glob.glob(os.path.join(folder, "gsec_valuation_*.xlsx")):
        iso = os.path.basename(f)[len("gsec_valuation_"):-len(".xlsx")]
        if os.path.exists(os.path.join(folder, f"tbill_{iso}.csv")):
            out.add(iso)
    return out


def list_dates(sid=None):
    """Dates with a G-Sec valuation file and T-bill rates: the study dates plus the dates entered through the app.
    On a shared deployment (sid given) only this visitor's entries are listed; on a laptop (sid None) all entries are."""
    out = set(_dates_in(MULTI))
    for iso in _dates_in(UPLOADED):
        if (sid is None and "~" not in iso) or (sid is not None and iso.endswith("~" + sid)):
            out.add(iso)
    return sorted(out, key=lambda k: (base(k), k))


def is_uploaded(iso):
    return iso not in _dates_in(MULTI) and iso in _dates_in(UPLOADED)


@dataclass
class Inputs:
    iso: str
    settle: date
    settle_table: pd.DataFrame
    bonds: pd.DataFrame        # all fixed-coupon G-Secs
    inp: pd.DataFrame          # FBIL input points
    tb: pd.DataFrame
    grid: np.ndarray           # tenors we publish (0.25-year steps)
    zcyc: pd.DataFrame = None  # FBIL published zero curve (optional: only for the comparison)
    par: pd.DataFrame = None   # FBIL published par curve (optional)

    @property
    def has_fbil(self):
        return self.zcyc is not None


@lru_cache(maxsize=32)
def load_inputs(iso):
    settle, table = md.detect_settlement(iso)
    g = md.load_gsec(iso)
    fixed = g[~g["is_frb"]].reset_index(drop=True)
    inp = fixed[fixed["is_input"]].reset_index(drop=True)
    zc = md.load_zcyc(iso) if os.path.exists(md._find(f"gsec_zcyc_and_strips_{iso}.xlsx")) else None
    try:
        par = md.load_par(iso)
    except Exception:
        par = None
    if zc is not None:
        grid = zc["tenor"].to_numpy(float)
    else:                       # FBIL publishes 0.25 to 50 years; stop one year after the last input bond
        last = max(md.yearfrac(m, settle) for m in inp["maturity"])
        grid = np.arange(0.25, min(50.0, np.floor((last + 1.0) * 4) / 4) + 1e-9, 0.25)
    return Inputs(iso, settle, table, fixed, inp, md.load_tbills(iso), grid, zc, par)


@lru_cache(maxsize=256)
def fit_curve(iso, method, lam, slope, use_3m):
    d = load_inputs(iso)
    if method == "exact":
        c, info = en.bootstrap(d.inp, d.tb, d.settle, use_3m=use_3m)
        return c, {"converged": bool(info["converged"]), "max_price_error": float(info["max_price_err"])}
    c, info = en.ls_fit(d.inp, d.tb, d.settle, lam=lam, lam_slope=slope, use_3m=use_3m)
    return c, {"converged": True, "mean_price_err_paise": float(info["mean_abs_price_err"] * 100)}


def clear_caches():
    load_inputs.cache_clear()
    fit_curve.cache_clear()


# ------------------------------------------------------------------ tables
def curve_table(iso, curve):
    """Our curve on the published tenors; FBIL's values and the gaps are filled in only when FBIL's files are loaded."""
    d = load_inputs(iso)
    g = d.grid
    z = curve.zero(g)
    t = pd.DataFrame({"tenor_years": g, "zero_semi_annual_pct": z, "zero_annualised_pct": en.semi_to_annual(z),
                      "par_semi_annual_pct": en.par_yield(curve, g), "discount_factor": curve.df(g)})
    if d.zcyc is not None:
        t["fbil_zero_semi_annual_pct"] = np.interp(g, d.zcyc["tenor"], d.zcyc["zero_semi"])
        t["zero_gap_bp"] = (t.zero_semi_annual_pct - t.fbil_zero_semi_annual_pct) * 100
    if d.par is not None:
        t["fbil_par_semi_annual_pct"] = np.interp(g, d.par["tenor"], d.par["par_semi"])
        t["par_gap_bp"] = (t.par_semi_annual_pct - t.fbil_par_semi_annual_pct) * 100
    return t


def summary(tab):
    """Gap metrics against FBIL (needs FBIL's files)."""
    tn = tab.tenor_years
    g = tab.zero_gap_bp.abs()
    out = {"All tenors": g.mean(), "Up to 14 years": g[tn <= 14].mean(), "Beyond 14 years": g[tn > 14].mean(),
           "Worst tenor": g.max(), "Tenors within 2 bp": (g <= 2).mean() * 100}
    out["Par curve, all tenors"] = tab.par_gap_bp.abs().mean() if "par_gap_bp" in tab else float("nan")
    return out


def level_summary(tab):
    """Curve levels, for dates without FBIL's curve to compare against."""
    pick = lambda col, t: float(tab.loc[(tab.tenor_years - t).abs().idxmin(), col])
    return {"2-year zero": pick("zero_semi_annual_pct", 2), "5-year zero": pick("zero_semi_annual_pct", 5),
            "10-year zero": pick("zero_semi_annual_pct", 10), "30-year zero": pick("zero_semi_annual_pct", 30),
            "10-year par": pick("par_semi_annual_pct", 10)}


def segment_table(tab):
    rows = []
    for a, b in SEGS:
        m = (tab.tenor_years > a) & (tab.tenor_years <= b)
        if not m.any():
            continue
        rows.append({"Segment (years)": f"{a}-{b}", "Mean absolute gap (bp)": tab.zero_gap_bp[m].abs().mean(),
                     "Mean gap, signed (bp)": tab.zero_gap_bp[m].mean(), "Largest gap (bp)": tab.zero_gap_bp[m].abs().max()})
    return pd.DataFrame(rows)


def bond_table(iso, curves):
    """Input bonds with the price error (model minus market dirty price, paise) under each fitted curve."""
    d = load_inputs(iso)
    flows, tgt = en.bond_flows(d.inp, d.settle)
    t = d.inp[["isin", "coupon", "maturity", "price", "ytm", "input_status"]].copy()
    t.insert(3, "years", [md.yearfrac(m, d.settle) for m in d.inp["maturity"]])
    for key, c in curves.items():
        model = np.array([np.dot(cf, c.df(tj)) for tj, cf in flows])
        t[f"{METHODS[key][0]}: price error (paise)"] = (model - tgt) * 100
    return t.rename(columns={"isin": "ISIN", "coupon": "Coupon %", "maturity": "Maturity", "price": "FBIL clean price",
                             "ytm": "FBIL YTM %", "input_status": "Traded / proxy", "years": "Years"})


def all_dates_table(lam, slope, use_3m, sid=None):
    """Comparison with FBIL on every date whose FBIL zero curve is available."""
    rows = []
    for iso in list_dates(sid):
        d = load_inputs(iso)
        if not d.has_fbil:
            continue
        r = {"Date": base(iso), "Role": "held out" if iso == HELD_OUT else ("tuning" if iso in md.DATES else "uploaded"),
             "Input bonds": len(d.inp)}
        for key in ("exact", "smooth"):
            sm = summary(curve_table(iso, fit_curve(iso, key, lam, slope, use_3m)[0]))
            r[f"{METHODS[key][0]}: all"] = sm["All tenors"]
            r[f"{METHODS[key][0]}: up to 14y"] = sm["Up to 14 years"]
            r[f"{METHODS[key][0]}: beyond 14y"] = sm["Beyond 14 years"]
        rows.append(r)
    return pd.DataFrame(rows)


# ------------------------------------------------------------------ SDL STRIPS
def sdl_path(iso):
    for p in (os.path.join(ROOT, "input_data", f"sdl_valuation_{iso}.xlsx"), os.path.join(MULTI, f"sdl_valuation_{iso}.xlsx"),
              os.path.join(UPLOADED, f"sdl_valuation_{iso}.xlsx")):
        if os.path.exists(p):
            return p
    return None


@lru_cache(maxsize=8)
def load_sdl(path):
    s = pd.read_excel(path, sheet_name="SDL", header=4)
    s.columns = ["isin", "description", "coupon", "maturity", "price", "ytm"]
    s = s.dropna(subset=["isin"]).copy()
    s["maturity"] = pd.to_datetime(s["maturity"]).dt.date
    s["strippable"] = [m.day == 2 and m.month in (1, 7) for m in s["maturity"]]
    return s.reset_index(drop=True)


def list_securities(iso, kind):
    """Securities that can be stripped. kind: 'gsec' (Government of India, from the G-Sec file) or 'sdl'."""
    d = load_inputs(iso)
    if kind == "sdl":
        p = sdl_path(iso)
        if p is None:
            return pd.DataFrame()
        s = load_sdl(p).copy()
        s["label"] = [f"{r.description}  ({r.isin}), matures {r.maturity}, FBIL YTM {r.ytm:.4f}%" for r in s.itertuples()]
    else:
        s = d.bonds[["isin", "coupon", "maturity", "price", "ytm"]].copy()
        s["description"] = [f"{r.coupon:.2f}% GS {r.maturity.year}" for r in s.itertuples()]
        s["strippable"] = [m.day == 2 and m.month in (1, 7) for m in s["maturity"]]
        s["label"] = [f"{r.description}  ({r.isin}), matures {r.maturity}, FBIL YTM {r.ytm:.4f}%" for r in s.itertuples()]
    return s[s["maturity"] > d.settle].sort_values("maturity", ascending=False).reset_index(drop=True)


def price_security(iso, kind, isin, curve_kind, lam, slope, use_3m):
    """STRIPS of one security under the RBI guidelines (para 15.2 normalisation). curve_kind: smooth | exact | fbil."""
    d = load_inputs(iso)
    sec = list_securities(iso, kind)
    b = sec[sec["isin"] == isin].iloc[0]
    if curve_kind == "fbil":
        if not d.has_fbil:
            raise ValueError("FBIL's ZCYC file is not loaded for this date")
        F = CubicSpline(d.zcyc["tenor"].to_numpy(float), d.zcyc["zero_semi"].to_numpy(float), bc_type="natural")
        zero_fn = lambda t: F(np.asarray(t, float))
    else:
        c = fit_curve(iso, curve_kind, lam, slope, use_3m)[0]
        zero_fn = lambda t: c.zero(np.asarray(t, float))
    accrued = float(md.accrued(b.coupon, b.maturity, d.settle))
    dirty = float(b.price) + accrued
    cfs = [(x, b.coupon / 2.0) for x, _ in md.cash_flows(b.coupon, b.maturity, d.settle)] + [(b.maturity, 100.0)]
    face = np.r_[np.full(len(cfs) - 1, b.coupon / 2.0), 100.0]
    st, factor = sp.price_strips(cfs, zero_fn, dirty, settle=d.settle, time_fn=lambda x: md.yearfrac(x, d.settle))
    st["strip_type"] = ["Coupon"] * (len(st) - 1) + ["Principal"]
    st["strip_name"] = [f"GS{x.strftime('%d%b%Y').upper()}C" for x in st["date"][:-1]] + [f"{b.coupon:.2f}%GS{b.maturity.strftime('%d%b%Y').upper()}P"]
    st["price_per_100_face"] = 100 * st["normalised_value"] / face
    st["implied_yield_semi_pct"] = sp.strip_yield_from_price(st["price_per_100_face"].to_numpy(), st["t_years"].to_numpy())
    info = {"description": b.description, "coupon": float(b.coupon), "maturity": b.maturity, "clean": float(b.price),
            "accrued": accrued, "dirty": dirty, "sum_pv": float(st.pv_unnormalised.sum()), "factor": float(factor),
            "n_coupon": len(st) - 1, "principal_price": float(st.price_per_100_face.iloc[-1]),
            "principal_yield": float(st.implied_yield_semi_pct.iloc[-1]),
            "principal_share_pct": float(st.normalised_value.iloc[-1] / dirty * 100),
            "years": float(st.t_years.iloc[-1]), "strippable": bool(b.strippable)}
    return st, info


price_sdl = lambda iso, isin, curve_kind, lam, slope, use_3m: price_security(iso, "sdl", isin, curve_kind, lam, slope, use_3m)


# ------------------------------------------------------------------ entering a new date
OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"     # start of an Excel 97-2003 (.xls) file


@lru_cache(maxsize=12)
def to_xlsx_bytes(data):
    """Excel 2007+ (.xlsx) files pass through unchanged. Excel 97-2003 (.xls) files, which is how FBIL publishes them,
    are converted to .xlsx in memory (same sheets, same cell positions), so the rest of the code reads one format.
    Raises ValueError with a plain explanation if the file cannot be used."""
    if data[:2] == b"PK":
        return data
    if data[:8] == OLE_MAGIC:
        try:
            import xlrd
        except ImportError as e:
            raise ValueError("FBIL's .xls files need the 'xlrd' package. Install it with  py -m pip install xlrd  "
                             "and restart the app (or save the file as .xlsx in Excel and upload that).") from e
        try:
            book = xlrd.open_workbook(file_contents=data)
        except Exception as e:
            raise ValueError(f"the .xls file could not be read ({e})") from e
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for sh in book.sheets():
            ws = wb.create_sheet(title=sh.name[:31])
            for r in range(sh.nrows):
                for c in range(sh.ncols):
                    t = sh.cell_type(r, c)
                    if t in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
                        continue
                    v = sh.cell_value(r, c)
                    if t == xlrd.XL_CELL_DATE:
                        try:
                            v = xlrd.xldate_as_datetime(v, book.datemode)
                        except Exception:
                            pass
                    elif t == xlrd.XL_CELL_BOOLEAN:
                        v = bool(v)
                    elif t == xlrd.XL_CELL_TEXT and v == "":
                        continue
                    ws.cell(row=r + 1, column=c + 1, value=v)
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()
    if data[:200].lstrip().lower().startswith(b"<"):
        raise ValueError("this file is HTML or XML with an .xls extension, not a real Excel workbook; open it in Excel, "
                         "use Save As > Excel Workbook (.xlsx), and upload that")
    raise ValueError("unrecognised file format; upload the Excel file (.xls or .xlsx) downloaded from FBIL")


def describe_file(data):
    """(date in the file header or None, error message or None), for showing next to an uploader."""
    try:
        return file_date(data, _raise=True), None
    except ValueError as e:
        return None, str(e)


def file_date(data, _raise=False):
    """Date printed in the header of an FBIL Excel file (.xlsx or .xls, cell B3), or None."""
    try:
        data = to_xlsx_bytes(data)
    except ValueError:
        if _raise:
            raise
        return None
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for sn in wb.sheetnames[:2]:
            for row in wb[sn].iter_rows(min_row=1, max_row=6, values_only=True):
                for c in row:
                    if isinstance(c, datetime):
                        return c.date()
                    if isinstance(c, str):
                        m = re.search(r"\b(\d{1,2})[-\s]([A-Za-z]{3})[A-Za-z]*[-\s](\d{4}|\d{2})\b", c)
                        if m:
                            dd, mon, yy = m.groups()
                            return datetime.strptime(f"{dd} {mon.title()} {yy}", "%d %b %Y" if len(yy) == 4 else "%d %b %y").date()
    except Exception:
        return None
    return None


def parse_tbills(text):
    """T-bill rates from pasted text. Accepts FBIL's table copied from the website ('11 Sep 2026  5:30:00 PM  7 Days  4.72')
    or plain lines such as '7 Days 4.72'. Returns {'7 Days': 4.72, '3 Months': 5.18, ...}."""
    out = {}
    for line in (text or "").splitlines():
        m = None
        for m in re.finditer(r"(\d+)\s*(days?|months?|years?|d|m|y)\b", line, re.I):
            pass
        if m is None:
            continue
        rate = re.search(r"\d+(?:\.\d+)?", line[m.end():])
        if rate is None or not (0 < float(rate.group()) < 25):
            continue
        n, unit = int(m.group(1)), m.group(2).lower()
        if unit.startswith("d"):
            label = f"{n} Days"
        elif unit.startswith("m"):
            label = "1 Month" if n == 1 else f"{n} Months"
        else:
            label = f"{12 * n} Months"
        out[label] = float(rate.group())
    return out


def save_upload(iso, gsec_bytes, tbill_rates, zcyc_bytes=None, sdl_bytes=None, sid=None):
    """Store one new FBIL date under input_data/uploaded/ using the project's file names and check that it works.
    Required: the G-Sec valuation file and the T-bill rates. Optional: the ZCYC and STRIPS file (adds the comparison with
    FBIL) and the SDL valuation file (adds SDL STRIPS). Raises ValueError, and removes the files again, if anything is wrong."""
    date.fromisoformat(iso)
    real_iso, key = iso, (f"{iso}~{sid}" if sid else iso)      # the key names the files and separates visitors
    if sid:
        purge_old_uploads()
    converted = []
    for name, data in (("G-Sec valuation", gsec_bytes), ("ZCYC and STRIPS", zcyc_bytes), ("SDL valuation", sdl_bytes)):
        try:
            converted.append(to_xlsx_bytes(data) if data else data)
        except ValueError as e:
            raise ValueError(f"The {name} file: {e}") from e
    gsec_bytes, zcyc_bytes, sdl_bytes = converted
    missing = [t for t in REQUIRED_TBILLS if t not in tbill_rates]
    if missing:
        raise ValueError("T-bill rates missing for: " + ", ".join(missing) + ". Paste FBIL's T-bill table or type lines like '7 Days 4.72'.")
    if real_iso in _dates_in(MULTI):
        raise ValueError(f"{real_iso} is part of the study data and cannot be replaced.")
    for name, data in (("G-Sec valuation", gsec_bytes), ("ZCYC and STRIPS", zcyc_bytes), ("SDL valuation", sdl_bytes)):
        if data:
            fd = file_date(data)
            if fd is not None and fd.isoformat() != real_iso:
                raise ValueError(f"The {name} file is dated {fd}, not {real_iso}. All files must be for the same date.")
    files = {f"gsec_valuation_{key}.xlsx": gsec_bytes}
    if zcyc_bytes:
        files[f"gsec_zcyc_and_strips_{key}.xlsx"] = zcyc_bytes
    if sdl_bytes:
        files[f"sdl_valuation_{key}.xlsx"] = sdl_bytes
    os.makedirs(UPLOADED, exist_ok=True)
    written = []
    try:
        for name, data in files.items():
            with open(os.path.join(UPLOADED, name), "wb") as fh:
                fh.write(data)
            written.append(name)
        pd.DataFrame({"tenor": list(tbill_rates), "rate_pct": [float(v) for v in tbill_rates.values()]}) \
            .to_csv(os.path.join(UPLOADED, f"tbill_{key}.csv"), index=False)
        written.append(f"tbill_{key}.csv")
        clear_caches()
        try:
            d = load_inputs(key)
        except ValueError as e:
            if "zero-size" in str(e):
                raise ValueError("no bond in the G-Sec file is still outstanding on that date; check that the date matches the file") from e
            raise
        except Exception as e:
            raise ValueError("the G-Sec valuation file could not be read (expected the FBIL workbook with sheets 'G-Sec' and 'Par Yield')") from e
        if len(d.inp) < 8:
            raise ValueError(f"only {len(d.inp)} 'Input Point' bonds found; this does not look like FBIL's G-Sec valuation file")
        best = float(d.settle_table["mean_paise"].min())
        if best > 2.0:
            raise ValueError(f"FBIL's clean prices cannot be reproduced from its YTMs at any settlement date (best fit {best:.1f} paise); "
                             "check that this is the G-Sec valuation file for that date")
        if zcyc_bytes and d.zcyc is None:
            raise ValueError("the ZCYC and STRIPS file could not be read (expected a sheet named 'ZCYC')")
        if sdl_bytes:
            load_sdl(sdl_path(key))
        for m in ("exact", "smooth"):
            fit_curve(key, m, *tuned_params(), True)
    except Exception as e:
        for name in written:
            try:
                os.remove(os.path.join(UPLOADED, name))
            except OSError:
                pass
        clear_caches()
        if isinstance(e, ValueError):
            raise ValueError(f"Could not use these files: {e}") from e
        raise ValueError(f"Could not use these files: {e}") from e
    return {"key": key, "settle": d.settle, "n_inputs": len(d.inp), "has_fbil": d.has_fbil, "has_sdl": bool(sdl_bytes)}


def delete_uploaded(iso):
    if not is_uploaded(iso):
        return
    for f in glob.glob(os.path.join(UPLOADED, f"*_{iso}.*")):
        os.remove(f)
    clear_caches()


# ------------------------------------------------------------------ figures
def _style():
    plt.rcParams.update({"font.size": 10.5, "axes.spines.top": False, "axes.spines.right": False,
                         "axes.grid": True, "grid.alpha": 0.25})


def fig_zero(iso, curves, show_bonds=True, xmax=50):
    _style()
    d = load_inputs(iso)
    g = d.grid
    two = d.has_fbil
    if two:
        fig, ax = plt.subplots(2, 1, figsize=(10, 6.6), gridspec_kw={"height_ratios": [2.2, 1]}, sharex=True)
        top = ax[0]
        fb = np.interp(g, d.zcyc["tenor"], d.zcyc["zero_semi"])
        top.plot(g, fb, color=NAVY, lw=2.4, label="FBIL published ZCYC")
    else:
        fig, top = plt.subplots(figsize=(10, 4.6)); ax = [top]
    gaps = []
    for key, c in curves.items():
        name, col, ls = METHODS[key]
        top.plot(g, c.zero(g), color=col if two else (NAVY if key == "smooth" else GREY), lw=1.8, ls=ls if two else "-", label=name)
        if two:
            gaps.append((c.zero(g) - fb) * 100)
            ax[1].plot(g, gaps[-1], color=col, lw=1.8, label=name)
    if show_bonds:
        top.scatter([md.yearfrac(m, d.settle) for m in d.inp["maturity"]], d.inp["ytm"], s=20, color=TEAL, zorder=5, label="Input bond YTMs")
    top.set_ylabel("Semi-annual zero rate (%)"); top.legend(frameon=False, loc="lower right", fontsize=9)
    top.set_title(f"Zero curve, {base(iso)}", loc="left", fontsize=11)
    if two:
        ax[1].axhspan(-0.5, 0.5, color=GREY, alpha=0.2); ax[1].axhline(0, color=GREY, lw=0.8); ax[1].axvline(14, color=NAVY, lw=0.8, ls="--")
        lo, hi = min(x.min() for x in gaps), max(x.max() for x in gaps)
        ax[1].set_ylim(min(lo - 2, -5), max(hi + 2, 5))
        ax[1].set_ylabel("Ours minus FBIL (bp)"); ax[1].set_xlabel("Tenor (years)"); ax[1].set_xlim(-0.5, xmax + 0.5)
    else:
        top.set_xlabel("Tenor (years)")
    plt.tight_layout()
    return fig


def fig_par(iso, curves):
    _style()
    d = load_inputs(iso)
    g = d.grid
    two = d.par is not None
    if two:
        fig, ax = plt.subplots(2, 1, figsize=(10, 6.0), gridspec_kw={"height_ratios": [2, 1]}, sharex=True)
        top = ax[0]
        fb = np.interp(g, d.par["tenor"], d.par["par_semi"])
        top.plot(g, fb, color=NAVY, lw=2.4, label="FBIL par yield")
    else:
        fig, top = plt.subplots(figsize=(10, 4.4)); ax = [top]
    gaps = []
    for key, c in curves.items():
        name, col, ls = METHODS[key]
        p = en.par_yield(c, g)
        top.plot(g, p, color=col if two else (NAVY if key == "smooth" else GREY), lw=1.8, ls=ls if two else "-", label=f"Par yield: {name.lower()}")
        if two:
            gaps.append((p - fb) * 100)
            ax[1].plot(g, gaps[-1], color=col, lw=1.6)
    top.set_ylabel("Par yield, semi-annual (%)"); top.legend(frameon=False, loc="lower right", fontsize=9)
    top.set_title(f"Par yield curve, {base(iso)}", loc="left", fontsize=11)
    if two:
        ax[1].axhspan(-0.5, 0.5, color=GREY, alpha=0.2); ax[1].axhline(0, color=GREY, lw=0.8)
        lo, hi = min(x.min() for x in gaps), max(x.max() for x in gaps)
        ax[1].set_ylim(min(lo - 1, -3), max(hi + 1, 3))
        ax[1].set_ylabel("Ours minus FBIL (bp)"); ax[1].set_xlabel("Tenor (years)")
    else:
        top.set_xlabel("Tenor (years)")
    plt.tight_layout()
    return fig


def fig_segments(seg_tables):
    _style()
    fig, ax = plt.subplots(figsize=(10, 4.4))
    n = len(seg_tables); w = 0.8 / n; x = np.arange(len(SEGS))
    top = max(t["Mean absolute gap (bp)"].max() for t in seg_tables.values()) * 1.15 or 1
    for j, (key, t) in enumerate(seg_tables.items()):
        name, col, _ = METHODS[key]
        xs = x + (j - (n - 1) / 2) * w
        ax.bar(xs, t["Mean absolute gap (bp)"], w, color=col, label=name)
        for xi, v in zip(xs, t["Mean absolute gap (bp)"]):
            ax.text(xi, v + top * 0.012, f"{v:.1f}", ha="center", fontsize=8.5)
    ax.axvline(4.5, color=NAVY, lw=0.8, ls="--")
    ax.set_xticks(x); ax.set_xticklabels([f"{a}–{b}" for a, b in SEGS]); ax.set_ylim(0, top)
    ax.set_xlabel("Maturity segment (years)"); ax.set_ylabel("Mean absolute gap to FBIL (bp)"); ax.legend(frameon=False)
    plt.tight_layout()
    return fig


def fig_dates(tab):
    _style()
    fig, ax = plt.subplots(1, 2, figsize=(11, 4.2))
    x = np.arange(len(tab)); w = 0.38
    for a, cols, ttl in zip(ax, ("all", "beyond 14y"), ("All tenors", "Tenors beyond 14 years")):
        for j, key in enumerate(("exact", "smooth")):
            name, col, _ = METHODS[key]
            a.bar(x + (j - 0.5) * w, tab[f"{name}: {cols}"], w, color=col, label=name)
        a.set_xticks(x); a.set_xticklabels([t[5:].replace("-", "/") for t in tab["Date"]])
        for i, r in enumerate(tab["Role"]):
            if r == "held out":
                a.axvspan(i - 0.5, i + 0.5, color=TEAL, alpha=0.1)
        a.set_title(ttl, loc="left", fontsize=11); a.set_ylabel("Mean absolute gap to FBIL (bp)")
    ax[0].legend(frameon=False, fontsize=9)
    plt.tight_layout()
    return fig


def fig_strips(st, info):
    _style()
    fig, ax = plt.subplots(1, 2, figsize=(11, 4.4))
    c = st[st.strip_type == "Coupon"]; p = st[st.strip_type == "Principal"]
    ax[0].bar(c.t_years, c.normalised_value, width=0.4, color=TEAL, label=f"Coupon STRIPS ({info['coupon'] / 2:.3f} face each)")
    ax[0].bar(p.t_years, p.normalised_value, width=0.4, color=ORANGE, label="Principal STRIP")
    ax[0].set_xlabel("Years to cash flow"); ax[0].set_ylabel("Normalised value per Rs 100 of the bond")
    ax[0].set_title("Value of each STRIP", loc="left", fontsize=11); ax[0].legend(frameon=False, fontsize=9)
    ax[1].plot(st.t_years, st.price_per_100_face, color=NAVY, marker="o", ms=3, lw=1.4)
    ax[1].set_xlabel("Years to maturity of the STRIP"); ax[1].set_ylabel("Price per Rs 100 face", color=NAVY)
    ax2 = ax[1].twinx(); ax2.plot(st.t_years, st.implied_yield_semi_pct, color=ORANGE, lw=1.6)
    ax2.set_ylabel("Implied yield, semi-annual (%)", color=ORANGE); ax2.grid(False); ax2.spines["right"].set_visible(True)
    ax[1].set_title("Price and implied yield", loc="left", fontsize=11)
    plt.tight_layout()
    return fig
